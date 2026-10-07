import datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch

OUT = Path(__file__).resolve().parent  # folder containing this script
START = dt.date(2026, 10, 8)
WEEKS = 9

TRACKS = {  # categorical palette, fixed order (validated)
    "Organization": "2A78D6",
    "Game engine": "EB6834",
    "Levels": "1BAF7A",
    "Bot": "EDA100",
    "Graphics": "E87BA4",
}

TASKS = [
    ("Organization", "Project requirements", 1, 1),
    ("Game engine", "Basics: movement, jump, collisions, obstacles, restart", 2, 3),
    ("Game engine", "Deterministic physics, headless mode", 2, 3),
    ("Levels", "Level format + 2–3 handmade levels", 2, 3),
    ("Graphics", "Graphic assets", 2, 8),
    ("Game engine", "Dash, climbing, stamina, buttons, checkpoints, coins", 4, 5),
    ("Bot", "Bot environment interface", 4, 4),
    ("Bot", "Bot v1", 4, 5),
    ("Levels", "Procedural generator + planner validation", 6, 7),
    ("Organization", "Human completion time measurement", 6, 6),
    ("Levels", "Test set of 10 levels", 7, 7),
    ("Bot", "Bot v2", 8, 8),
    ("Bot", "Evaluation on 10 test levels", 8, 9),
    ("Game engine", "Performance optimization", 9, 9),
    ("Organization", "Final report", 9, 9),
]

# week -> (label, fill) — colors match the schedule sheet
MILESTONES = {
    1: ("Project requirements", "6AA84F"),
    3: ("Report", None),
    5: ("Milestone I", "E07B39"),
    7: ("Report", None),
    9: ("Deadline", "E81010"),
}

# ---------------- XLSX ----------------
wb = Workbook()
ws = wb.active
ws.title = "Gantt"
F = "Arial"
thin = Side(style="thin", color="D0D0D0")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

HDR = 4
first_col = 6  # F
headers = ["Track", "Task", "Start (week)", "End (week)", "Duration (weeks)"]
for i, h in enumerate(headers, 1):
    c = ws.cell(row=HDR, column=i, value=h)
for i, h in enumerate(headers, 1):
    for r in (HDR, HDR + 1, HDR + 2):
        ws.cell(row=r, column=i).border = box
    c = ws.cell(row=HDR, column=i)
    c.font = Font(name=F, bold=True)
    c.alignment = center
    ws.merge_cells(start_row=HDR, start_column=i, end_row=HDR + 2, end_column=i)

for w in range(1, WEEKS + 1):
    col = first_col + w - 1
    L = ws.cell(row=HDR, column=col).column_letter
    ws.cell(row=HDR, column=col, value=w)                       # week number
    if w == 1:
        ws.cell(row=HDR + 1, column=col, value=START)           # start date
        ws.cell(row=HDR + 1, column=col).comment = Comment(
            "Date of the first schedule deadline (08.10.2026).", "Claude")
    else:
        prev = ws.cell(row=HDR + 1, column=col - 1).column_letter
        ws.cell(row=HDR + 1, column=col, value=f"={prev}{HDR + 1}+7")
    ws.cell(row=HDR + 1, column=col).number_format = "dd.mm"
    label, fill = MILESTONES.get(w, ("", None))
    ws.cell(row=HDR + 2, column=col, value=label)
    for r in (HDR, HDR + 1, HDR + 2):
        c = ws.cell(row=r, column=col)
        c.font = Font(name=F, bold=(r != HDR + 2), size=9 if r == HDR + 2 else 10,
                      color="FFFFFF" if (r == HDR + 2 and fill) else "000000")
        c.alignment = center
        c.border = box
    if fill:
        ws.cell(row=HDR + 2, column=col).fill = PatternFill("solid", fgColor=fill)
    ws.column_dimensions[L].width = 11
ws.cell(row=HDR + 1, column=first_col).font = Font(name=F, bold=True, color="0000FF")

first = HDR + 3
last = first + len(TASKS) - 1
for i, (track, name, s, e) in enumerate(TASKS):
    r = first + i
    ws.cell(row=r, column=1, value=track)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=s)
    ws.cell(row=r, column=4, value=e)
    ws.cell(row=r, column=5, value=f"=D{r}-C{r}+1")
    for col in range(1, first_col + WEEKS):
        c = ws.cell(row=r, column=col)
        c.border = box
        c.font = Font(name=F, color="0000FF" if col in (3, 4) else "000000")
        c.alignment = Alignment(vertical="center", horizontal="left" if col <= 2 else "center",
                                wrap_text=(col == 2))
    ws.row_dimensions[r].height = 30
    # track color bar in column A
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=TRACKS[track])
    ws.cell(row=r, column=1).font = Font(name=F, bold=True, color="FFFFFF")

fL = ws.cell(row=1, column=first_col).column_letter
lL = ws.cell(row=1, column=first_col + WEEKS - 1).column_letter
rng = f"{fL}{first}:{lL}{last}"
for track, color in TRACKS.items():
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND({fL}${HDR}>=$C{first},{fL}${HDR}<=$D{first},$A{first}="{track}")'],
        fill=PatternFill("solid", fgColor=color, bgColor=color), stopIfTrue=True))

ws.column_dimensions["A"].width = 13
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 9
ws.column_dimensions["D"].width = 9
ws.column_dimensions["E"].width = 9
ws.row_dimensions[HDR + 2].height = 30
ws.freeze_panes = ws.cell(row=first, column=first_col)

# legend
lr = last + 2
ws.cell(row=lr, column=1, value="Legend").font = Font(name=F, bold=True)
for i, (track, color) in enumerate(TRACKS.items()):
    c = ws.cell(row=lr + 1 + i, column=1, value=track)
    c.fill = PatternFill("solid", fgColor=color)
    c.font = Font(name=F, bold=True, color="FFFFFF")
ws.cell(row=lr + 1, column=2, value="Blue text = editable cells").font = Font(name=F, color="0000FF")
ws.cell(row=lr + 2, column=2, value="Plan drafted with Claude, 07.10.2026").font = Font(name=F, italic=True, color="52514E")

ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT / "gantt-chart.xlsx")

# ---------------- PNG (16:9 slide) ----------------
plt.rcParams["font.family"] = "Arial"
fig, ax = plt.subplots(figsize=(16, 9), dpi=120)
fig.patch.set_alpha(0)
ax.set_facecolor("none")
n = len(TASKS)
for i, (track, name, s, e) in enumerate(TASKS):
    y = n - 1 - i
    ax.barh(y, e - s + 1, left=s - 0.5, height=0.62, color="#" + TRACKS[track],
            edgecolor="none")
ax.set_yticks(range(n))
ax.set_yticklabels([t[1] for t in reversed(TASKS)], fontsize=12, color="#0b0b0b")
dates = [START + dt.timedelta(weeks=w) for w in range(WEEKS)]
ax.set_xticks(range(1, WEEKS + 1))
ax.set_xticklabels([f"Week {w}\n{d:%d.%m}" for w, d in zip(range(1, WEEKS + 1), dates)],
                   fontsize=11, color="#52514e")
ax.set_xlim(0.5, WEEKS + 0.65)
ax.set_ylim(-0.7, n - 0.3)
for w in range(1, WEEKS + 2):
    ax.axvline(w - 0.5, color="#e6e5e1", linewidth=1, zorder=0)
for w, (label, fill) in MILESTONES.items():
    col = "#" + fill if fill else "#898781"
    ax.axvline(w + 0.5, color=col, linewidth=2, linestyle="--", zorder=3)
    ax.text(w + 0.5, n - 0.1, label, ha="center", va="bottom", fontsize=11,
            fontweight="bold", color="#0b0b0b",
            bbox=dict(boxstyle="round,pad=0.3", fc=col if fill else "#f0efec",
                      ec="none", alpha=0.9 if fill else 1))
for sp in ("top", "right", "left"):
    ax.spines[sp].set_visible(False)
ax.spines["bottom"].set_color("#898781")
ax.tick_params(axis="y", length=0)
ax.legend(handles=[Patch(color="#" + c, label=t) for t, c in TRACKS.items()],
          loc="lower center", bbox_to_anchor=(0.5, -0.16), ncol=len(TRACKS),
          frameon=False, fontsize=12)
fig.tight_layout()
fig.savefig(OUT / "gantt-chart.png", transparent=True)
print("ok")
